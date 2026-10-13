"""Map every premium field for one line: where it is written, what feeds it, and
how it rolls up to the account total.

    python scripts/premium/01_premium_map.py --line GL
    python scripts/premium/01_premium_map.py --line CF

Writes <LINE>_Algorithm/<LINE>_Premium_Map.xlsx -- Summary, RollUpTree,
PremiumSources, PremiumFields. Read-only against the corpus; exit 0 always."""
import re, os, glob, collections, sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import argparse
LINES = {
    "GL": dict(pkg_glob=r"F:/ISO_ERC_Files/General_Liability/*/GL CW *",
               root="GeneralLiability", out=os.path.join("GL_Algorithm", "GL_Premium_Map.xlsx"),
               dict_=os.path.join("GL_Algorithm", "GL_Data_Dictionary.xlsx")),
    "CF": dict(pkg_glob=r"F:/ISO_ERC_Files/CF/*/CFCW*",
               root="CommercialProperty", out=os.path.join("CF_Algorithm", "CF_Premium_Map.xlsx"),
               dict_=os.path.join("CF_Algorithm", "CF_Data_Dictionary.xlsx")),
}
_ap = argparse.ArgumentParser()
_ap.add_argument("--line", choices=sorted(LINES), default="GL")
_A = _ap.parse_args()
_C = LINES[_A.line]
REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PKG = sorted(glob.glob(_C["pkg_glob"]))[-1]
OUT = os.path.join(REPO, _C["out"])
DICT = os.path.join(REPO, _C["dict_"])
ROOT = _C["root"]


def rule_body(text, name):
    i = text.find('Name="%s"' % name)
    if i < 0:
        return ""
    j = text.find("</rul:Rule>", i)
    return text[i:j] if j > 0 else text[i:]


nodes = {}          # group -> record
for path in sorted(glob.glob(os.path.join(PKG, "Rules", "*.xml"))):
    t = open(path, encoding="utf-8", errors="replace").read()
    fname = os.path.basename(path)

    calc = rule_body(t, "CalculateTotalPremium")
    if not calc:
        continue
    grp = re.search(r'DataDefGroup="([^"]+)"', calc[:300])
    grp = grp.group(1) if grp else fname.replace("Rules.Rule.xml", "")

    aggregates = "ToDataDef=\"ErcCalculatedTotalPremium\"" in calc
    # Repeating children are summed with ForEach over an <X>Table; singleton children
    # with Locate on the group itself. Reading only the first misses 59 groups.
    # NOTE the digits in the character class: group names like
    # ...DispositionOfAct2002 exist, and [A-Za-z]+ silently drops them.
    rep = set(re.findall(r'<rul:ForEach AtDataDef="([A-Za-z0-9]+)Table/', calc))
    sing = set(re.findall(r'<rul:Locate AtDataDef="([A-Za-z0-9]+)"', calc))
    children = sorted(rep | sing)
    # a leaf returns its own Premium under a PremiumIndicator gate
    own = 'FromDataDef="Premium"' in calc

    sp = rule_body(t, "SetPremium")
    inputs, lookups, rules_called = [], [], []
    if sp:
        inputs = sorted(set(re.findall(r'FromDataDef="([A-Za-z/\.]+)"', sp)))
        lookups = sorted(set(re.findall(r'MatrixDef="([A-Za-z]+)Def"', sp)))
        rules_called = sorted(set(re.findall(r'Rule="([A-Za-z]+)"', sp)))

    spi = rule_body(t, "SetPremiumIndicator")

    nodes[grp] = {
        "group": grp, "file": fname,
        "aggregates": aggregates, "children": children, "own_premium": own,
        "n_repeating": len(rep), "n_singleton": len(sing),
        "has_setpremium": bool(sp), "setpremium_chars": len(sp),
        "inputs": inputs, "lookups": lookups, "rules_called": rules_called,
        "has_indicator": bool(spi),
    }

# ---- depth from the root -----------------------------------------------------
depth, parent = {ROOT: 0}, {}
queue = collections.deque([ROOT])
while queue:
    g = queue.popleft()
    for c in nodes.get(g, {}).get("children", []):
        if c not in depth:
            depth[c] = depth[g] + 1
            parent[c] = g
            queue.append(c)

orphans = [g for g in nodes if g not in depth]


def path_to_root(g):
    out, seen = [], set()
    while g in parent and g not in seen:
        seen.add(g)
        out.append(parent[g])
        g = parent[g]
    return " < ".join(out)


# ---- premium field census ----------------------------------------------------
wbd = openpyxl.load_workbook(DICT, read_only=True)
declared = collections.Counter()
for r in wbd["Fields"].iter_rows(min_row=2, values_only=True):
    if r[1] and "Premium" in r[1]:
        declared[r[1]] += 1

written = collections.Counter()
writers = collections.defaultdict(set)
for path in glob.glob(os.path.join(PKG, "Rules", "*.xml")):
    t = open(path, encoding="utf-8", errors="replace").read()
    for m in re.finditer(r'ToDataDef="([A-Za-z]*Premium[A-Za-z]*)"', t):
        written[m.group(1)] += 1
        writers[m.group(1)].add(os.path.basename(path).replace("Rules.Rule.xml", ""))

# ---- workbook ----------------------------------------------------------------
wb = openpyxl.Workbook()


def sheet(name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i in range(1, len(headers) + 1):
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), ws.max_row)
    return ws


tree_rows = []
for g in sorted(nodes, key=lambda x: (depth.get(x, 99), x)):
    n = nodes[g]
    tree_rows.append([
        g, depth.get(g, ""), parent.get(g, "(root)" if g == ROOT else "(unreached)"),
        "aggregator" if n["aggregates"] else "leaf",
        len(n["children"]), n["n_repeating"], n["n_singleton"],
        ", ".join(n["children"][:5]) + (" …" if len(n["children"]) > 5 else ""),
        "yes" if n["own_premium"] else "", "yes" if n["has_setpremium"] else "",
        "yes" if n["has_indicator"] else "", path_to_root(g)[:220], n["file"],
    ])
sheet("RollUpTree",
      ["DataDefGroup", "Depth from account", "Rolls up into", "Role", "Children summed",
       "  repeating (ForEach)", "  singleton (Locate)", "Children (first 5)",
       "Contributes own Premium", "Has SetPremium", "Has PremiumIndicator",
       "Path to account total", "Rule file"],
      tree_rows, [58, 17, 46, 13, 16, 20, 20, 60, 22, 15, 20, 70, 52])

src_rows = []
for g in sorted(nodes):
    n = nodes[g]
    if not n["has_setpremium"]:
        continue
    src_rows.append([
        g, depth.get(g, ""), n["setpremium_chars"],
        len(n["inputs"]), ", ".join(x for x in n["inputs"] if not x.startswith("../"))[:300],
        ", ".join(x for x in n["inputs"] if x.startswith("../"))[:200],
        ", ".join(n["lookups"])[:200], ", ".join(n["rules_called"])[:200], n["file"],
    ])
sheet("PremiumSources",
      ["DataDefGroup", "Depth", "SetPremium size (chars)", "N inputs", "Own-record fields read",
       "Ancestor fields read", "Rate tables read", "Sub-rules called", "Rule file"],
      src_rows, [58, 8, 20, 10, 70, 50, 46, 46, 52])

fld_rows = []
for f in sorted(set(declared) | set(written)):
    fld_rows.append([
        f, declared.get(f, 0), written.get(f, 0),
        "computed" if written.get(f) else "input / carried",
        ", ".join(sorted(writers.get(f, []))[:4])[:180],
    ])
sheet("PremiumFields",
      ["Field", "Groups declaring it", "Rule writes", "Kind", "Written by (first 4)"],
      fld_rows, [46, 20, 13, 18, 72])

summary = [
    ["Line", _A.line],
    ["Source package", os.path.basename(PKG)],
    ["Account total field", ROOT + ".ErcCalculatedTotalPremium"],
    ["", ""],
    ["Distinct premium-named fields in the DataDef", len(declared)],
    ["  of which any rule computes", len([f for f in declared if written.get(f)])],
    ["  carried but never computed by a rule", len([f for f in declared if not written.get(f)])],
    ["", ""],
    ["Groups with a CalculateTotalPremium", len(nodes)],
    ["  aggregators (write ErcCalculatedTotalPremium)", sum(1 for n in nodes.values() if n["aggregates"])],
    ["  leaves (return own Premium)", sum(1 for n in nodes.values() if not n["aggregates"])],
    ["Groups with a SetPremium (compute a premium)", sum(1 for n in nodes.values() if n["has_setpremium"])],
    ["Deepest level below the account total", max(d for d in depth.values())],
    ["Groups never reached from the account root", len(orphans)],
]
ws = wb.create_sheet("Summary", 0)
for r in summary:
    ws.append(r)
ws.cell(row=1, column=1).font = Font(bold=True)
ws.column_dimensions["A"].width = 46
ws.column_dimensions["B"].width = 52

del wb["Sheet"]
wb.save(OUT)

print("package:", os.path.basename(PKG))
print("groups with CalculateTotalPremium:", len(nodes))
print("  aggregators:", sum(1 for n in nodes.values() if n["aggregates"]),
      "| leaves:", sum(1 for n in nodes.values() if not n["aggregates"]))
print("  with SetPremium:", sum(1 for n in nodes.values() if n["has_setpremium"]))
print("max depth:", max(depth.values()), "| unreached groups:", len(orphans))
print("depth histogram:", dict(sorted(collections.Counter(depth.values()).items())))
print("premium fields: declared", len(declared), "| computed", len([f for f in declared if written.get(f)]))
print("->", OUT)
if orphans[:6]:
    print("unreached examples:", orphans[:6])
